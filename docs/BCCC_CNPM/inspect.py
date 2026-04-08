import openpyxl
import os

filepath = r'D:\Hutech\Congnghephanmem\TERENCE\TERENCE-main\docs\BCCC_CNPM\TERENCE_TestExecution_HUTECH_Pro.xlsx'
wb = openpyxl.load_workbook(filepath)
print("Sheets:", wb.sheetnames)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n--- {sheet} ---")
    headers = []
    
    # Try finding the row with recognizable headers like "Test Case ID" or "Status"
    header_row = 1
    for row in range(1, 10):
        vals = [str(ws.cell(row=row, column=c).value) for c in range(1, ws.max_column + 1)]
        if 'Test Case ID' in vals or 'Module' in vals or 'Trạng thái' in vals or 'Final Status' in vals or 'BugID' in vals:
            header_row = row
            headers = vals
            break
            
    print(f"Header row: {header_row}")
    print("Headers:", [(i+1, h) for i, h in enumerate(headers) if h != 'None'])
